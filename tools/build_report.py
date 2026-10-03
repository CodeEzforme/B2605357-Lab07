from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    Image as RLImage,
    KeepTogether,
    NextPageTemplate,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
SCREENS = ROOT / "assets" / "report"
DOCX_OUT = ROOT / "BaoCaoKyNangSo.docx"
PDF_OUT = ROOT / "BaoCaoKyNangSo.pdf"

NAVY = "0B1F33"
CYAN = "12B8C4"
PALE = "EAF6F7"
INK = "17212B"
MUTED = "5D6B78"
LINE = "D9E1E5"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def set_run_font(run, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def style_paragraph_runs(paragraph):
    for run in paragraph.runs:
        set_run_font(run)


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = "Bài thực hành 07    •    Báo cáo Kỹ năng số"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = "Footer"
    for run in p.runs:
        set_run_font(run)
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def add_title_block(doc, number, title, intro):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(number)
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(CYAN)
    h = doc.add_heading(title, level=1)
    h.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph(intro)
    p.paragraph_format.space_after = Pt(13)
    p.style = "Intro"


def add_caption(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(7)
    for r in p.runs:
        set_run_font(r)
        r.italic = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_docx_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r)
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(9)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cells[i], top=90, bottom=90)
            if row_index % 2:
                set_cell_shading(cells[i], "F3F8FA")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    set_run_font(r)
                    r.font.size = Pt(8.7)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_docx_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(8)
    for name, size in (("Heading 1", 19), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
    intro = styles.add_style("Intro", 1)
    intro.font.name = "Arial"
    intro.font.size = Pt(11)
    intro.font.color.rgb = RGBColor.from_string(MUTED)
    intro.paragraph_format.line_spacing = 1.2


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    configure_docx_styles(doc)
    add_footer(section)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(38)
    r = p.add_run("BÀI THỰC HÀNH 07")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(CYAN)
    p = doc.add_paragraph()
    r = p.add_run("07")
    set_run_font(r)
    r.bold = True
    r.font.size = Pt(68)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p.paragraph_format.space_after = Pt(0)
    title = doc.add_paragraph(style="Title")
    title.add_run("Báo cáo Dịch vụ tiện ích và sử dụng sáng tạo")
    subtitle = doc.add_paragraph("Google Forms  •  GitHub Pages  •  Quản lý công việc  •  Quản lý dự án")
    subtitle.paragraph_format.space_after = Pt(36)
    for r in subtitle.runs:
        set_run_font(r)
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor.from_string(MUTED)
    add_docx_table(doc, ["Nội dung", "Thông tin"], [
        ("Họ và tên", "Trần Ngọc Mơ"),
        ("Mã sinh viên", "B2605357"),
        ("Lớp và khóa", "26D1A2 - Khóa 52"),
        ("Ngành", "Truyền thông đa phương tiện"),
        ("Thời gian thực hiện", "Tháng 10 năm 2026"),
        ("Ngày lập báo cáo", "03/10/2026"),
    ], [5.0, 11.8])
    p = doc.add_paragraph("Báo cáo tổng hợp cấu hình biểu mẫu, website cá nhân, kế hoạch công việc tháng 10/2026 và kế hoạch dự án nhóm. Các tệp nguồn đi kèm cho phép tái tạo tài nguyên trên Google và GitHub.")
    p.paragraph_format.space_before = Pt(20)

    doc.add_page_break()
    add_title_block(doc, "BÀI 1", "Tiện ích Google Forms", "Hai biểu mẫu được đặc tả đầy đủ và có mã Google Apps Script để tạo tự động trong tài khoản Google.")
    add_docx_table(doc, ["Biểu mẫu", "Trường thông tin", "Điều kiện tự động"], [
        ("Thông tin sinh viên khóa 52", "Mã SV, họ tên, phái, ngày sinh, nơi sinh, ngành, mã lớp, cố vấn, email, điện thoại, thông tin cha mẹ, địa chỉ gia đình", "Đóng ngày 01/12/2026 theo múi giờ Asia/Bangkok"),
        ("Đăng ký tham quan TMA", "Mã SV, họ tên, phái, ngày sinh, nơi sinh, ngành, mã lớp, cố vấn, email, điện thoại", "Đóng ngay khi đủ 80 phản hồi"),
    ], [4.2, 8.4, 4.2])
    doc.add_heading("Cách triển khai", level=2)
    add_bullets(doc, [
        "Mở script.google.com, tạo dự án và dán nội dung automation/Code.gs.",
        "Chạy hàm buildLab07 một lần để tạo hai form, hai trigger và lịch tháng 10/2026.",
        "Chạy getLab07Links để lấy link điền và link chỉnh sửa, sau đó lưu link điền vào README.md.",
        "Mở hai link trong cửa sổ ẩn danh để kiểm tra quyền truy cập trước khi nộp.",
    ])
    doc.add_page_break()
    add_title_block(doc, "BÀI 1", "Biểu mẫu thông tin sinh viên đầy đủ", "Ảnh chụp toàn bộ biểu mẫu từ tiêu đề đến nút Gửi.")
    picture = doc.add_picture(str(SCREENS / "form-sinhvien-day-du.png"), width=Cm(8.2))
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Hình 1a  Toàn bộ Google Form thông tin sinh viên khóa 52")

    doc.add_page_break()
    add_title_block(doc, "BÀI 1", "Biểu mẫu đăng ký tham quan đầy đủ", "Ảnh chụp toàn bộ biểu mẫu từ tiêu đề đến nút Gửi.")
    picture = doc.add_picture(str(SCREENS / "form-tma-day-du.png"), width=Cm(12.3))
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, "Hình 1b  Toàn bộ Google Form đăng ký tham quan Công ty TMA")

    doc.add_page_break()
    add_title_block(doc, "BÀI 2", "Website giới thiệu bản thân", "Website của Trần Ngọc Mơ được xây dựng bằng HTML, CSS và xuất bản công khai bằng GitHub Pages.")
    doc.add_picture(str(SCREENS / "website.png"), width=Cm(16.9))
    add_caption(doc, "Hình 2  Giao diện đầu trang của website hồ sơ sinh viên")
    add_docx_table(doc, ["Khu vực", "Nội dung"], [
        ("Giới thiệu", "Trần Ngọc Mơ, sinh ngày 18/09/2007, đến từ Bạc Liêu"),
        ("Quá trình học", "Sinh viên khóa 52, lớp 26D1A2, mã sinh viên B2605357"),
        ("Lĩnh vực chuyên môn", "Truyền thông đa phương tiện, nội dung số và công cụ cộng tác"),
        ("Kinh nghiệm và sản phẩm", "Google Forms, Google Calendar, website cá nhân và kế hoạch dự án Lab 07"),
    ], [5.0, 11.8])
    p = doc.add_paragraph("Đường dẫn công khai: https://codeezforme.github.io/B2605357-Lab07/")
    for r in p.runs:
        set_run_font(r)
        r.font.color.rgb = RGBColor.from_string(CYAN)
        r.bold = True

    doc.add_page_break()
    add_title_block(doc, "BÀI 3", "Kế hoạch công việc tháng 10 năm 2026", "Kế hoạch được tổ chức theo tuần, kết hợp ghi chú, danh sách việc cần làm và lịch có nhắc hẹn.")
    doc.add_picture(str(SCREENS / "calendar-google.png"), width=Cm(16.9))
    add_caption(doc, "Hình 3  Google Calendar thực tế với kế hoạch tháng 10 năm 2026")
    add_docx_table(doc, ["Tuần", "Trọng tâm", "Kết quả cần đạt"], [
        ("01–04/10", "Ổn định lịch học và xác định ưu tiên", "Kế hoạch tháng và danh sách việc"),
        ("05–11/10", "Lập trình và bài tập nhóm", "Cập nhật tiến độ tuần"),
        ("12–18/10", "Hoàn thiện bài tập đang làm", "Bản nộp trước hạn"),
        ("19–25/10", "Ôn giữa kỳ và tích hợp bài nhóm", "Kiểm tra chất lượng bài"),
        ("26–31/10", "Hoàn thiện Lab 07 và tổng kết", "Bộ bài nộp và kế hoạch tháng mới"),
    ], [3.2, 7.0, 6.6])
    p = doc.add_paragraph("Tệp calendar/KeHoachThang10-2026.ics có thể nhập trực tiếp vào Google Calendar. Script cũng tạo một lịch riêng và nhắc trước 30 phút cho từng sự kiện.")

    doc.add_page_break()
    add_title_block(doc, "BÀI 4", "Kế hoạch dự án nhóm", "Nhóm chọn Dự án 4 trong danh sách mẫu và lập kế hoạch phát triển website quản lý chi tiêu cá nhân theo nhiều vòng lặp.")
    doc.add_picture(str(SCREENS / "kanban.png"), width=Cm(16.9))
    add_caption(doc, "Hình 4  Bảng Kanban kế hoạch và phân công dự án")
    add_docx_table(doc, ["Thành viên", "Trách nhiệm chính", "Sản phẩm phụ trách"], [
        ("Trần Ngọc Mơ", "Điều phối, giao diện và tài liệu", "Phạm vi, giao diện chính, biểu đồ và hướng dẫn"),
        ("Thành viên 2", "Dữ liệu và xử lý nghiệp vụ", "Cấu trúc dữ liệu, giao dịch và bộ lọc"),
        ("Cả nhóm", "Kiểm thử và rà soát", "Tiêu chí kiểm thử, tích hợp và demo"),
    ], [3.3, 6.2, 7.3])

    doc.add_page_break()
    add_title_block(doc, "PHỤ LỤC", "Cấu trúc tệp bài làm", "Các tệp được tổ chức để dễ kiểm tra, triển khai và cập nhật trước khi nộp.")
    add_docx_table(doc, ["Tệp hoặc thư mục", "Mục đích"], [
        ("README.md", "Lưu link hai Google Forms và link website GitHub Pages"),
        ("automation/", "Mã Apps Script và hướng dẫn tạo tài nguyên Google"),
        ("site/", "Mã nguồn website cá nhân"),
        ("calendar/", "Tệp lịch tháng 10/2026 để nhập Google Calendar"),
        ("project/", "Dữ liệu kế hoạch dự án dùng cho Trello"),
        ("BaoCaoKyNangSo.docx", "Bản Word nguồn của báo cáo"),
        ("BaoCaoKyNangSo.pdf", "Tệp báo cáo nộp chính"),
    ], [6.0, 10.8])
    doc.add_heading("Trạng thái hoàn tất", level=2)
    add_bullets(doc, [
        "Hai link Google Forms công khai đã được lưu trong README.md.",
        "Website cá nhân đã dùng thông tin thật và được xuất bản bằng GitHub Pages.",
        "Kế hoạch tháng 10/2026 đã được tạo và chụp từ Google Calendar.",
        "Bài 4 chỉ lập kế hoạch và phân công, đúng phạm vi yêu cầu của đề.",
        "Báo cáo PDF đã được xuất và kiểm tra trực quan trước khi nộp.",
    ])

    for paragraph in doc.paragraphs:
        style_paragraph_runs(paragraph)
    doc.core_properties.title = "Báo cáo Dịch vụ tiện ích và sử dụng sáng tạo"
    doc.core_properties.subject = "Bài thực hành 07"
    doc.core_properties.author = "Trần Ngọc Mơ"
    doc.save(DOCX_OUT)


def register_pdf_fonts():
    regular = Path(r"C:\Windows\Fonts\arial.ttf")
    bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
    italic = Path(r"C:\Windows\Fonts\ariali.ttf")
    pdfmetrics.registerFont(TTFont("Arial", str(regular)))
    pdfmetrics.registerFont(TTFont("Arial-Bold", str(bold)))
    pdfmetrics.registerFont(TTFont("Arial-Italic", str(italic)))
    pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial-Italic")


def scaled_image(path, max_width, max_height):
    with Image.open(path) as im:
        width, height = im.size
    scale = min(max_width / width, max_height / height)
    return RLImage(str(path), width=width * scale, height=height * scale)


def pdf_table(headers, rows, widths):
    data = [[Paragraph(h, PDF_STYLES["table_header"]) for h in headers]]
    for row in rows:
        data.append([Paragraph(str(c), PDF_STYLES["table_cell"]) for c in row])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#" + LINE)),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]
    for i in range(1, len(data)):
        if i % 2 == 0:
            commands.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F3F8FA")))
    table.setStyle(TableStyle(commands))
    return table


def pdf_bullets(items):
    return [Paragraph("• " + item, PDF_STYLES["body"]) for item in items]


def draw_page(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D9E1E5"))
    canvas.line(18 * mm, 14 * mm, 192 * mm, 14 * mm)
    canvas.setFont("Arial", 8)
    canvas.setFillColor(colors.HexColor("#" + MUTED))
    canvas.drawString(18 * mm, 9 * mm, "Bài thực hành 07  •  Báo cáo Kỹ năng số")
    canvas.drawRightString(192 * mm, 9 * mm, str(doc.page))
    canvas.restoreState()


def build_pdf_styles():
    global PDF_STYLES
    base = getSampleStyleSheet()
    PDF_STYLES = {
        "eyebrow": ParagraphStyle("eyebrow", parent=base["Normal"], fontName="Arial-Bold", fontSize=8, leading=10, textColor=colors.HexColor("#" + CYAN), spaceAfter=4),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="Arial-Bold", fontSize=20, leading=23, textColor=colors.black, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Arial-Bold", fontSize=13, leading=16, textColor=colors.black, spaceBefore=8, spaceAfter=6),
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName="Arial", fontSize=9.5, leading=13.5, textColor=colors.HexColor("#" + INK), spaceAfter=6),
        "intro": ParagraphStyle("intro", parent=base["BodyText"], fontName="Arial", fontSize=10.5, leading=15, textColor=colors.HexColor("#" + MUTED), spaceAfter=12),
        "caption": ParagraphStyle("caption", parent=base["Normal"], fontName="Arial-Italic", fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.HexColor("#" + MUTED), spaceBefore=5, spaceAfter=9),
        "table_header": ParagraphStyle("table_header", parent=base["Normal"], fontName="Arial-Bold", fontSize=8, leading=10, alignment=TA_CENTER, textColor=colors.white),
        "table_cell": ParagraphStyle("table_cell", parent=base["Normal"], fontName="Arial", fontSize=8, leading=10.5, textColor=colors.HexColor("#" + INK)),
    }


def build_pdf():
    register_pdf_fonts()
    build_pdf_styles()
    frame = Frame(18 * mm, 17 * mm, 174 * mm, 262 * mm, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    template = PageTemplate(id="main", frames=[frame], onPage=draw_page)
    pdf = BaseDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=17 * mm, title="Báo cáo Dịch vụ tiện ích và sử dụng sáng tạo", author="Trần Ngọc Mơ")
    pdf.addPageTemplates([template])
    story = []

    story += [Spacer(1, 12 * mm), Paragraph("BÀI THỰC HÀNH 07", PDF_STYLES["eyebrow"])]
    big = ParagraphStyle("big", fontName="Arial-Bold", fontSize=62, leading=62, textColor=colors.HexColor("#" + NAVY), spaceAfter=2)
    title = ParagraphStyle("cover_title", fontName="Arial-Bold", fontSize=27, leading=31, textColor=colors.black, spaceAfter=10)
    sub = ParagraphStyle("sub", fontName="Arial", fontSize=11, leading=15, textColor=colors.HexColor("#" + MUTED), spaceAfter=22)
    story += [Paragraph("07", big), Paragraph("Báo cáo Dịch vụ tiện ích<br/>và sử dụng sáng tạo", title), Paragraph("Google Forms  •  GitHub Pages  •  Quản lý công việc  •  Quản lý dự án", sub)]
    story += [pdf_table(["Nội dung", "Thông tin"], [
        ("Họ và tên", "Trần Ngọc Mơ"),
        ("Mã sinh viên", "B2605357"),
        ("Lớp và khóa", "26D1A2 - Khóa 52"),
        ("Ngành", "Truyền thông đa phương tiện"),
        ("Thời gian thực hiện", "Tháng 10 năm 2026"),
        ("Ngày lập báo cáo", "03/10/2026"),
    ], [48 * mm, 126 * mm]), Spacer(1, 10 * mm), Paragraph("Báo cáo tổng hợp cấu hình biểu mẫu, website cá nhân, kế hoạch công việc tháng 10/2026 và kế hoạch dự án nhóm. Các tệp nguồn đi kèm cho phép tái tạo tài nguyên trên Google và GitHub.", PDF_STYLES["intro"])]

    story += [PageBreak(), Paragraph("BÀI 1", PDF_STYLES["eyebrow"]), Paragraph("Tiện ích Google Forms", PDF_STYLES["h1"]), Paragraph("Hai biểu mẫu được đặc tả đầy đủ và có mã Google Apps Script để tạo tự động trong tài khoản Google.", PDF_STYLES["intro"])]
    story += [pdf_table(["Biểu mẫu", "Trường thông tin", "Điều kiện tự động"], [
        ("Thông tin sinh viên khóa 52", "Mã SV, họ tên, phái, ngày sinh, nơi sinh, ngành, mã lớp, cố vấn, email, điện thoại, thông tin cha mẹ, địa chỉ gia đình", "Đóng ngày 01/12/2026 theo múi giờ Asia/Bangkok"),
        ("Đăng ký tham quan TMA", "Mã SV, họ tên, phái, ngày sinh, nơi sinh, ngành, mã lớp, cố vấn, email, điện thoại", "Đóng ngay khi đủ 80 phản hồi"),
    ], [42 * mm, 85 * mm, 47 * mm]), Spacer(1, 4 * mm), Paragraph("Cách triển khai", PDF_STYLES["h2"])]
    story += pdf_bullets([
        "Mở script.google.com, tạo dự án và dán nội dung automation/Code.gs.",
        "Chạy hàm buildLab07 một lần để tạo hai form, hai trigger và lịch tháng 10/2026.",
        "Chạy getLab07Links để lấy link điền và link chỉnh sửa, sau đó lưu link điền vào README.md.",
        "Mở hai link trong cửa sổ ẩn danh để kiểm tra quyền truy cập trước khi nộp.",
    ])
    story += [
        PageBreak(),
        Paragraph("BÀI 1", PDF_STYLES["eyebrow"]),
        Paragraph("Biểu mẫu thông tin sinh viên đầy đủ", PDF_STYLES["h1"]),
        Paragraph("Ảnh chụp toàn bộ biểu mẫu từ tiêu đề đến nút Gửi.", PDF_STYLES["intro"]),
        scaled_image(SCREENS / "form-sinhvien-day-du.png", 174 * mm, 205 * mm),
        Paragraph("Hình 1a  Toàn bộ Google Form thông tin sinh viên khóa 52", PDF_STYLES["caption"]),
        PageBreak(),
        Paragraph("BÀI 1", PDF_STYLES["eyebrow"]),
        Paragraph("Biểu mẫu đăng ký tham quan đầy đủ", PDF_STYLES["h1"]),
        Paragraph("Ảnh chụp toàn bộ biểu mẫu từ tiêu đề đến nút Gửi.", PDF_STYLES["intro"]),
        scaled_image(SCREENS / "form-tma-day-du.png", 174 * mm, 205 * mm),
        Paragraph("Hình 1b  Toàn bộ Google Form đăng ký tham quan Công ty TMA", PDF_STYLES["caption"]),
    ]

    story += [PageBreak(), Paragraph("BÀI 2", PDF_STYLES["eyebrow"]), Paragraph("Website giới thiệu bản thân", PDF_STYLES["h1"]), Paragraph("Website của Trần Ngọc Mơ được xây dựng bằng HTML, CSS và xuất bản công khai bằng GitHub Pages.", PDF_STYLES["intro"]), scaled_image(SCREENS / "website.png", 174 * mm, 132 * mm), Paragraph("Hình 2  Giao diện đầu trang của website hồ sơ sinh viên", PDF_STYLES["caption"])]
    story += [pdf_table(["Khu vực", "Nội dung"], [
        ("Giới thiệu", "Trần Ngọc Mơ, sinh ngày 18/09/2007, đến từ Bạc Liêu"),
        ("Quá trình học", "Sinh viên khóa 52, lớp 26D1A2, mã sinh viên B2605357"),
        ("Lĩnh vực chuyên môn", "Truyền thông đa phương tiện, nội dung số và công cụ cộng tác"),
        ("Kinh nghiệm và sản phẩm", "Google Forms, Google Calendar, website cá nhân và kế hoạch dự án Lab 07"),
    ], [48 * mm, 126 * mm]), Spacer(1, 3 * mm), Paragraph("Đường dẫn công khai: <font color='#12B8C4'><b>https://codeezforme.github.io/B2605357-Lab07/</b></font>", PDF_STYLES["body"])]

    story += [PageBreak(), Paragraph("BÀI 3", PDF_STYLES["eyebrow"]), Paragraph("Kế hoạch công việc tháng 10 năm 2026", PDF_STYLES["h1"]), Paragraph("Kế hoạch được tổ chức theo tuần, kết hợp ghi chú, danh sách việc cần làm và lịch có nhắc hẹn.", PDF_STYLES["intro"]), scaled_image(SCREENS / "calendar-google.png", 174 * mm, 122 * mm), Paragraph("Hình 3  Google Calendar thực tế với kế hoạch tháng 10 năm 2026", PDF_STYLES["caption"])]
    story += [pdf_table(["Tuần", "Trọng tâm", "Kết quả cần đạt"], [
        ("01–04/10", "Ổn định lịch học và xác định ưu tiên", "Kế hoạch tháng và danh sách việc"),
        ("05–11/10", "Lập trình và bài tập nhóm", "Cập nhật tiến độ tuần"),
        ("12–18/10", "Hoàn thiện bài tập đang làm", "Bản nộp trước hạn"),
        ("19–25/10", "Ôn giữa kỳ và tích hợp bài nhóm", "Kiểm tra chất lượng bài"),
        ("26–31/10", "Hoàn thiện Lab 07 và tổng kết", "Bộ bài nộp và kế hoạch tháng mới"),
    ], [32 * mm, 76 * mm, 66 * mm]), Spacer(1, 3 * mm), Paragraph("Tệp calendar/KeHoachThang10-2026.ics có thể nhập trực tiếp vào Google Calendar. Script cũng tạo một lịch riêng và nhắc trước 30 phút cho từng sự kiện.", PDF_STYLES["body"])]

    story += [PageBreak(), Paragraph("BÀI 4", PDF_STYLES["eyebrow"]), Paragraph("Kế hoạch dự án nhóm", PDF_STYLES["h1"]), Paragraph("Nhóm chọn Dự án 4 trong danh sách mẫu và lập kế hoạch phát triển website quản lý chi tiêu cá nhân theo nhiều vòng lặp.", PDF_STYLES["intro"]), scaled_image(SCREENS / "kanban.png", 174 * mm, 108 * mm), Paragraph("Hình 4  Bảng Kanban kế hoạch và phân công dự án", PDF_STYLES["caption"])]
    story += [pdf_table(["Thành viên", "Trách nhiệm chính", "Sản phẩm phụ trách"], [
        ("Trần Ngọc Mơ", "Điều phối, giao diện và tài liệu", "Phạm vi, giao diện chính, biểu đồ và hướng dẫn"),
        ("Thành viên 2", "Dữ liệu và xử lý nghiệp vụ", "Cấu trúc dữ liệu, giao dịch và bộ lọc"),
        ("Cả nhóm", "Kiểm thử và rà soát", "Tiêu chí kiểm thử, tích hợp và demo"),
    ], [35 * mm, 62 * mm, 77 * mm])]

    story += [PageBreak(), Paragraph("PHỤ LỤC", PDF_STYLES["eyebrow"]), Paragraph("Cấu trúc tệp bài làm", PDF_STYLES["h1"]), Paragraph("Các tệp được tổ chức để dễ kiểm tra, triển khai và cập nhật trước khi nộp.", PDF_STYLES["intro"])]
    story += [pdf_table(["Tệp hoặc thư mục", "Mục đích"], [
        ("README.md", "Lưu link hai Google Forms và link website GitHub Pages"),
        ("automation/", "Mã Apps Script và hướng dẫn tạo tài nguyên Google"),
        ("site/", "Mã nguồn website cá nhân"),
        ("calendar/", "Tệp lịch tháng 10/2026 để nhập Google Calendar"),
        ("project/", "Dữ liệu kế hoạch dự án dùng cho Trello"),
        ("BaoCaoKyNangSo.docx", "Bản Word nguồn của báo cáo"),
        ("BaoCaoKyNangSo.pdf", "Tệp báo cáo nộp chính"),
    ], [60 * mm, 114 * mm]), Spacer(1, 5 * mm), Paragraph("Trạng thái hoàn tất", PDF_STYLES["h2"])]
    story += pdf_bullets([
        "Hai link Google Forms công khai đã được lưu trong README.md.",
        "Website cá nhân đã dùng thông tin thật và được xuất bản bằng GitHub Pages.",
        "Kế hoạch tháng 10/2026 đã được tạo và chụp từ Google Calendar.",
        "Bài 4 chỉ lập kế hoạch và phân công, đúng phạm vi yêu cầu của đề.",
        "Báo cáo PDF đã được xuất và kiểm tra trực quan trước khi nộp.",
    ])
    pdf.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)
